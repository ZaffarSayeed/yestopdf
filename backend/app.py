from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import uuid
import zipfile
from pathlib import Path
from typing import Annotated

import fitz
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt
from openpyxl import Workbook
import pdfplumber

try:
    import pikepdf  # type: ignore
except Exception:
    pikepdf = None

try:
    from pdf2docx import Converter as PDF2DOCX  # type: ignore
except Exception:
    PDF2DOCX = None

APP_DIR = Path(__file__).resolve().parent
TMP_ROOT = Path(os.getenv("PDFTOOLS_TMP", APP_DIR / "tmp")).resolve()
TMP_ROOT.mkdir(parents=True, exist_ok=True)
MAX_UPLOAD_MB = int(os.getenv("MAX_UPLOAD_MB", "100"))
MAX_UPLOAD_BYTES = MAX_UPLOAD_MB * 1024 * 1024

app = FastAPI(title="YesToPDF API", version="1.0.0")
# Vite can select another port when 5173 is already occupied, and browsers treat
# localhost and 127.0.0.1 as different origins. Keep an explicit configurable list
# for deployments, plus a localhost-only regex for local development.
_default_origins = (
    "http://localhost:5173,http://127.0.0.1:5173,"
    "http://localhost:5174,http://127.0.0.1:5174,"
    "http://localhost:3000,http://127.0.0.1:3000,"
    "http://localhost:8000,http://127.0.0.1:8000"
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[x.strip().rstrip("/") for x in os.getenv("CORS_ORIGINS", _default_origins).split(",") if x.strip()],
    allow_origin_regex=r"https?://(localhost|127\.0\.0\.1)(:\d+)?",
    allow_credentials=False,
    allow_methods=["POST", "GET", "OPTIONS"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

ALLOWED = {
    "pdf-to-word": {"pdf"}, "word-to-pdf": {"docx", "doc"},
    "pdf-to-jpg": {"pdf"}, "pdf-to-png": {"pdf"}, "jpg-to-pdf": {"jpg", "jpeg", "png", "tif", "tiff"},
    "pdf-to-excel": {"pdf"}, "pdf-to-ppt": {"pdf"}, "rotate": {"pdf"},
    "compress": {"pdf"}, "merge": {"pdf"}, "split": {"pdf"},
    "protect": {"pdf"}, "unlock": {"pdf"}, "sign": {"pdf"},
}


def cleanup(path: Path):
    shutil.rmtree(path, ignore_errors=True)


def ext(name: str) -> str:
    return Path(name or "").suffix.lower().lstrip(".")


def safe_stem(name: str) -> str:
    stem = Path(name or "document").stem
    return re.sub(r"[^A-Za-z0-9._-]+", "_", stem)[:80] or "document"


async def save_upload(upload: UploadFile, folder: Path, allowed: set[str]) -> Path:
    e = ext(upload.filename or "")
    if e not in allowed:
        raise HTTPException(415, f"Unsupported file type. Expected: {', '.join(sorted(allowed))}")
    out = folder / f"input.{e}"
    size = 0
    with out.open("wb") as f:
        while True:
            chunk = await upload.read(1024 * 1024)
            if not chunk:
                break
            size += len(chunk)
            if size > MAX_UPLOAD_BYTES:
                raise HTTPException(413, f"File exceeds {MAX_UPLOAD_MB} MB limit")
            f.write(chunk)
    return out


def response(path: Path, filename: str, background: BackgroundTasks, media_type: str = "application/octet-stream"):
    data = path.read_bytes()
    folder = path.parent
    background.add_task(cleanup, folder)
    return StreamingResponse(io.BytesIO(data), media_type=media_type, headers={
        "Content-Disposition": f'attachment; filename="{filename}"',
        "Cache-Control": "no-store",
        "X-Content-Type-Options": "nosniff",
    })


def pdf_to_word(src: Path, out: Path):
    if PDF2DOCX is not None:
        cv = PDF2DOCX(str(src))
        try:
            cv.convert(str(out))
        finally:
            cv.close()
        return
    # Editable fallback: reconstruct text blocks and embedded raster images page-by-page.
    pdf = fitz.open(str(src))
    doc = Document()
    for i, page in enumerate(pdf):
        if i:
            doc.add_page_break()
        rect = page.rect
        section = doc.sections[-1]
        section.page_width = DocxInches(rect.width / 72)
        section.page_height = DocxInches(rect.height / 72)
        blocks = page.get_text("dict").get("blocks", [])
        text_blocks = [b for b in blocks if b.get("type") == 0]
        text_blocks.sort(key=lambda b: (b["bbox"][1], b["bbox"][0]))
        if not text_blocks:
            p = doc.add_paragraph("[Scanned/image-only page: OCR engine required for editable text]")
            p.runs[0].italic = True
        for b in text_blocks:
            lines = b.get("lines", [])
            text = "\n".join("".join(s.get("text", "") for s in line.get("spans", [])) for line in lines).strip()
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = DocxPt(2)
                p.add_run(text)
        for b in blocks:
            if b.get("type") == 1 and b.get("image"):
                try:
                    img = b["image"]
                    stream = io.BytesIO(img)
                    width = max(1.0, min(6.5, (b["bbox"][2] - b["bbox"][0]) / 72))
                    doc.add_picture(stream, width=DocxInches(width))
                except Exception:
                    pass
    doc.save(str(out))


def word_to_pdf(src: Path, out: Path):
    # Best available Windows engine first, then LibreOffice.
    try:
        import win32com.client  # type: ignore
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(src.resolve()), ReadOnly=True)
        doc.ExportAsFixedFormat(str(out.resolve()), 17)
        doc.Close(False)
        word.Quit()
        return
    except Exception:
        pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        tmp = out.parent / "lo"
        tmp.mkdir(exist_ok=True)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)], check=True, timeout=180)
        produced = tmp / (src.stem + ".pdf")
        if produced.exists():
            shutil.move(str(produced), str(out)); return
    raise RuntimeError("Word-to-PDF requires Microsoft Word or LibreOffice on the server.")


def pdf_images(src: Path, out: Path, fmt: str):
    pdf = fitz.open(str(src))
    files = []
    for i, page in enumerate(pdf, 1):
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        p = out.parent / f"page-{i}.{fmt.lower()}"
        pix.save(str(p)); files.append(p)
    if len(files) == 1:
        shutil.copy2(files[0], out.parent / f"{out.stem}.{fmt.lower()}")
        out.unlink(missing_ok=True)
        return
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for p in files: z.write(p, p.name)


def images_to_pdf(src: Path, out: Path):
    # src is a zip containing normalized images for multi-file input.
    images = []
    if src.suffix == ".zip":
        folder = src.parent / "imgs"
        folder.mkdir(exist_ok=True)
        # Do not extract archive paths: read only image members and write safe basenames.
        with zipfile.ZipFile(src) as z:
            for idx, info in enumerate(z.infolist(), 1):
                if info.is_dir() or Path(info.filename).suffix.lower() not in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}:
                    continue
                if info.file_size > MAX_UPLOAD_BYTES:
                    continue
                safe_name = f"image-{idx}{Path(info.filename).suffix.lower()}"
                target = folder / safe_name
                with z.open(info) as source, target.open("wb") as dest:
                    shutil.copyfileobj(source, dest)
        images = sorted([p for p in folder.iterdir() if p.suffix.lower() in {".jpg", ".jpeg", ".png", ".tif", ".tiff"}])
    else:
        images = [src]
    if not images: raise RuntimeError("No supported images found")
    pil = []
    for p in images:
        im = Image.open(p).convert("RGB"); pil.append(im)
    pil[0].save(out, save_all=True, append_images=pil[1:], resolution=150)


def pdf_to_excel(src: Path, out: Path):
    wb = Workbook(); ws = wb.active; ws.title = "Page 1"
    wrote = False
    with pdfplumber.open(str(src)) as pdf:
        for pi, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables()
            if pi > 1:
                ws = wb.create_sheet(f"Page {pi}")
            row = 1
            if tables:
                for table in tables:
                    for vals in table:
                        for col, val in enumerate(vals or [], 1): ws.cell(row=row, column=col, value=val or "")
                        row += 1
                    row += 1
                wrote = True
            else:
                text = page.extract_text() or ""
                for line in text.splitlines():
                    ws.cell(row=row, column=1, value=line); row += 1
                wrote = wrote or bool(text)
    if not wrote:
        ws["A1"] = "No extractable table/text found. OCR may be required for scanned PDFs."
    wb.save(out)


def pdf_to_ppt(src: Path, out: Path):
    prs = Presentation()
    pdf = fitz.open(str(src))
    for i, page in enumerate(pdf):
        pw, ph = page.rect.width, page.rect.height
        prs.slide_width = Inches(pw / 72)
        prs.slide_height = Inches(ph / 72)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # Visual fallback/background first: this preserves the original page appearance.
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        stream = io.BytesIO(pix.tobytes("png"))
        slide.shapes.add_picture(stream, 0, 0, width=prs.slide_width, height=prs.slide_height)
        sx, sy = prs.slide_width / pw, prs.slide_height / ph
        # Editable text is placed above the visual background.
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            text = "\n".join("".join(s.get("text", "") for s in line.get("spans", [])) for line in block.get("lines", [])).strip()
            if not text:
                continue
            x0, y0, x1, y1 = block["bbox"]
            box = slide.shapes.add_textbox(int(x0 * sx), int(y0 * sy), max(1, int((x1-x0) * sx)), max(1, int((y1-y0) * sy)))
            tf = box.text_frame; tf.clear(); tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = text
            try:
                span = block["lines"][0]["spans"][0]
                p.font.size = Pt(float(span.get("size", 12)))
            except Exception:
                pass
    prs.save(out)


def basic_pdf_operation(src: Path, out: Path, op: str, **kwargs):
    pdf = fitz.open(str(src))
    if op == "merge":
        raise RuntimeError("merge uses dedicated handler")
    if op == "rotate":
        angle = int(kwargs.get("angle", 90)) % 360
        for p in pdf: p.set_rotation((p.rotation + angle) % 360)
        pdf.save(str(out), garbage=4, deflate=True)
    elif op == "compress":
        pdf.save(str(out), garbage=4, clean=True, deflate=True, deflate_images=True, deflate_fonts=True)
    elif op == "protect":
        pw = kwargs.get("password") or ""
        if not pw: raise ValueError("Password is required")
        pdf.save(str(out), garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw=pw, owner_pw=pw)
    elif op == "unlock":
        pw = kwargs.get("password") or ""
        if not pdf.authenticate(pw): raise ValueError("Incorrect PDF password")
        pdf.save(str(out), garbage=4, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
    elif op == "sign":
        text = kwargs.get("signature") or "Signed with YesToPDF"
        page = pdf[-1]
        page.insert_text((72, page.rect.height - 60), text, fontsize=14, color=(0.1,0.1,0.1))
        pdf.save(str(out), garbage=4, deflate=True)
    else: raise RuntimeError("Unknown PDF operation")


def merge_pdfs(files: list[Path], out: Path):
    dest = fitz.open()
    for p in files:
        src = fitz.open(str(p)); dest.insert_pdf(src); src.close()
    dest.save(str(out), garbage=4, deflate=True)


def split_pdf(src: Path, out: Path, start: int, end: int):
    pdf = fitz.open(str(src)); total = len(pdf)
    start = max(1, start); end = min(total, end)
    if start > end: raise ValueError(f"Invalid page range. PDF has {total} pages.")
    dest = fitz.open(); dest.insert_pdf(pdf, from_page=start-1, to_page=end-1); dest.save(str(out), garbage=4, deflate=True)


@app.get("/health")
@app.get("/api/health")
def health():
    return {"ok": True, "version": app.version, "engines": {"pdf2docx": PDF2DOCX is not None, "pikepdf": pikepdf is not None, "word_com": os.name == "nt", "libreoffice": bool(shutil.which("soffice") or shutil.which("libreoffice"))}}


@app.post("/api/process/{tool}")
async def process(tool: str, background: BackgroundTasks, files: Annotated[list[UploadFile], File(...)], password: str = Form(""), angle: int = Form(90), start_page: int = Form(1), end_page: int = Form(1), signature: str = Form("")):
    if tool not in ALLOWED: raise HTTPException(404, "Unknown tool")
    if not files: raise HTTPException(400, "No files uploaded")
    job = TMP_ROOT / uuid.uuid4().hex; job.mkdir()
    try:
        inputs = []
        for f in files:
            p = await save_upload(f, job, ALLOWED[tool])
            inputs.append((p, f.filename or p.name))
        src = inputs[0][0]
        stem = safe_stem(inputs[0][1])
        out = job / "result"
        filename = "result.bin"; media = "application/octet-stream"
        if tool == "pdf-to-word": out = job/(stem+".docx"); pdf_to_word(src,out); filename=out.name; media="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif tool == "word-to-pdf": out=job/(stem+".pdf"); word_to_pdf(src,out); filename=out.name; media="application/pdf"
        elif tool in {"pdf-to-jpg","pdf-to-png"}:
            fmt = "JPG" if tool.endswith("jpg") else "PNG"
            out = job / (stem + f"-{fmt.lower()}.zip")
            pdf_images(src, out, fmt)
            if len(fitz.open(str(src))) == 1:
                out = job / f"{out.stem}.{fmt.lower()}"
                filename = out.name
                media = "image/jpeg" if fmt == "JPG" else "image/png"
            else:
                filename = out.name
                media = "application/zip"
        elif tool == "jpg-to-pdf":
            if len(inputs) > 1:
                z=job/"images.zip"; 
                with zipfile.ZipFile(z,"w") as zz:
                    for p,n in inputs: zz.write(p,n)
                src=z
            out=job/(stem+".pdf"); images_to_pdf(src,out); filename=out.name; media="application/pdf"
        elif tool == "pdf-to-excel": out=job/(stem+".xlsx"); pdf_to_excel(src,out); filename=out.name; media="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif tool == "pdf-to-ppt": out=job/(stem+".pptx"); pdf_to_ppt(src,out); filename=out.name; media="application/vnd.openxmlformats-officedocument.presentationml.presentation"
        elif tool in {"rotate","compress","protect","unlock","sign"}:
            suffix = {"rotate":"-rotated","compress":"-compressed","protect":"-protected","unlock":"-unlocked","sign":"-signed"}[tool]
            out=job/(stem+suffix+".pdf"); basic_pdf_operation(src,out,tool,password=password,angle=angle,signature=signature); filename=out.name; media="application/pdf"
        elif tool == "merge":
            out=job/(stem+"-merged.pdf"); merge_pdfs([p for p,_ in inputs],out); filename=out.name; media="application/pdf"
        elif tool == "split":
            out=job/(stem+f"-pages-{start_page}-{end_page}.pdf"); split_pdf(src,out,start_page,end_page); filename=out.name; media="application/pdf"
        return response(out, filename, background, media)
    except HTTPException:
        cleanup(job); raise
    except Exception as e:
        cleanup(job)
        raise HTTPException(422, str(e))
